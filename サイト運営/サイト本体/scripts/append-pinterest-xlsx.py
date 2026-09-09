#!/usr/bin/env python3
"""
generate-pinterest-pins.js が書き出した _pending-manifest-rows.json を、
pinterest-pins-お金/pinterest-お金.xlsx の続きの行に追記する。

xlsxは「投稿日・投稿時間(21:00〜21:30 / 7:00〜7:30 の交互スケジュール)」を
ユーザーが手動管理する運用のため、CSVではなくこの形式で直接追記する
(2026-09-10、姉妹サイト〔美容〕のExcel運用に合わせて変更)。

列構成(1行目ヘッダー): A投稿日(数式) / B投稿時間 / C画像名(slug) / D記事URL /
E ステータス(既定値「未投稿」) / F カテゴリ名 / G ボード名 / H ピンタイトル案 / I ピン説明文案

使い方:
  python scripts/append-pinterest-xlsx.py
  (サイト運営\サイト本体 がカレントディレクトリ。事前に
   node scripts/generate-pinterest-pins.js <slug>... を実行しておくこと)
"""
import json
import os
import sys
from copy import copy

try:
    import openpyxl
except ImportError:
    print("openpyxlが見つかりません。`pip install openpyxl` を実行してください。", file=sys.stderr)
    sys.exit(1)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(SCRIPT_DIR, "..", "..", "..", "pinterest-pins-お金"))
XLSX_PATH = os.path.join(OUT_DIR, "pinterest-お金.xlsx")
PENDING_PATH = os.path.join(OUT_DIR, "_pending-manifest-rows.json")

# テンプレート行: 偶数行(21:00〜21:30)はrow2、奇数行(7:00〜7:30)はrow3のスタイルを流用する。
TEMPLATE_EVEN = 2
TEMPLATE_ODD = 3
COLS = ["A", "B", "C", "D", "E", "F", "G", "H", "I"]


def main():
    if not os.path.exists(PENDING_PATH):
        print(f"[append] {PENDING_PATH} が見つかりません。先にgenerate-pinterest-pins.jsを実行してください。", file=sys.stderr)
        sys.exit(1)
    if not os.path.exists(XLSX_PATH):
        print(f"[append] {XLSX_PATH} が見つかりません。xlsx本体を用意してから実行してください。", file=sys.stderr)
        sys.exit(1)

    with open(PENDING_PATH, "r", encoding="utf-8") as f:
        rows = json.load(f)
    if not rows:
        print("[append] 追記対象が0件のため何もしません。")
        return

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    # 画像名(C列)が入っている最終行を探し、その次の行から追記する。
    last_row = 1
    for r in range(2, ws.max_row + 1):
        if ws[f"C{r}"].value not in (None, ""):
            last_row = r
    start_row = last_row + 1

    existing_slugs = {
        ws[f"C{r}"].value for r in range(2, last_row + 1) if ws[f"C{r}"].value
    }

    appended = 0
    for i, row in enumerate(rows):
        slug = row["slug"]
        if slug in existing_slugs:
            print(f"[append] 既に登録済みのためスキップ: {slug}")
            continue
        r = start_row + appended
        is_even = (r - 2) % 2 == 0
        template_row = TEMPLATE_EVEN if is_even else TEMPLATE_ODD
        time_slot = "21:00〜21:30" if is_even else "7:00〜7:30"
        values = {
            "A": "=DATE(2026,9,11)+INT((ROW()-ROW($A$2))/2)",
            "B": time_slot,
            "C": slug,
            "D": row["url"],
            "E": "未投稿",
            "F": row["category"],
            "G": row["board"],
            "H": row["pinTitle"],
            "I": row["pinDescription"],
        }
        for col in COLS:
            src = ws[f"{col}{template_row}"]
            dst = ws[f"{col}{r}"]
            dst.value = values[col]
            dst._style = copy(src._style)
        appended += 1
        print(f"[append] 追記: row{r} {slug}")

    wb.save(XLSX_PATH)
    os.remove(PENDING_PATH)
    print(f"[append] 完了: {appended}件追記 / {XLSX_PATH}")


if __name__ == "__main__":
    main()
